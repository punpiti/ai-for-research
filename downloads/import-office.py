#!/usr/bin/env python3
"""Extract PPTX or XLSX structure into readable Markdown."""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, timezone
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any, Iterable


PACKAGE_GROUPS = {
    "pptx": ("python-pptx",),
    "xlsx": ("openpyxl",),
}


def markdown_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    return str(value).replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def package_versions(kind: str) -> dict[str, str]:
    found: dict[str, str] = {}
    for package in PACKAGE_GROUPS[kind]:
        try:
            found[package] = version(package)
        except PackageNotFoundError:
            found[package] = "missing"
    return found


def markdown_table(rows: Iterable[Iterable[Any]]) -> list[str]:
    materialized = [list(row) for row in rows]
    width = max((len(row) for row in materialized), default=0)
    if width == 0:
        return []
    output = [
        "| " + " | ".join(f"Column {index}" for index in range(1, width + 1)) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in materialized:
        padded = row + [""] * (width - len(row))
        output.append("| " + " | ".join(markdown_cell(value) for value in padded) + " |")
    output.append("")
    return output


def convert_pptx(source: Path) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(source)
    output: list[str] = []
    for slide_number, slide in enumerate(deck.slides, start=1):
        output.extend((f"## Slide {slide_number}", ""))
        title = slide.shapes.title
        title_shape_id = title.shape_id if title is not None else None
        if title is not None and title.has_text_frame and title.text.strip():
            output.extend((f"### {title.text.strip()}", ""))

        shapes = sorted(slide.shapes, key=lambda shape: (shape.top, shape.left))
        for shape in shapes:
            if shape.shape_id == title_shape_id:
                continue
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                output.extend(markdown_table(rows))
                continue
            if getattr(shape, "has_chart", False):
                output.extend(("_[Chart or plotted data: inspect the original slide.]_", ""))
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                output.extend(("_[Image: inspect the original slide.]_", ""))
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            paragraphs = [paragraph.text.strip() for paragraph in shape.text_frame.paragraphs]
            paragraphs = [paragraph for paragraph in paragraphs if paragraph]
            for paragraph in paragraphs:
                output.append(f"- {paragraph}")
            if paragraphs:
                output.append("")
        notes = []
        if slide.has_notes_slide:
            notes = [paragraph.text.strip() for paragraph in slide.notes_slide.notes_text_frame.paragraphs]
        notes = [paragraph for paragraph in notes if paragraph]
        if notes:
            output.extend(("#### Speaker notes", ""))
            output.extend(f"- {paragraph}" for paragraph in notes)
            output.append("")
    return "\n".join(output).strip() + "\n"


def convert_xlsx(source: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=False)
    output: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            output.extend((f"## Worksheet: {worksheet.title}", ""))
            if worksheet.sheet_state != "visible":
                output.extend((f"_Sheet state: {worksheet.sheet_state}._", ""))
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            while rows and all(value is None for value in rows[-1]):
                rows.pop()
            width = max(
                (index for row in rows for index, value in enumerate(row, start=1) if value is not None),
                default=0,
            )
            output.extend(markdown_table(row[:width] for row in rows) if width else ["_Worksheet is empty._", ""])
    finally:
        workbook.close()
    return "\n".join(output).strip() + "\n"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    parser.add_argument("--kind", required=True, choices=sorted(PACKAGE_GROUPS))
    parser.add_argument("--receipt", type=Path)
    args = parser.parse_args()

    text = convert_pptx(args.source) if args.kind == "pptx" else convert_xlsx(args.source)
    if not text.strip():
        text = "_No extractable text was found. Compare this file with the original._\n"
    args.target.write_text(text, encoding="utf-8")

    versions = package_versions(args.kind)
    if args.receipt:
        args.receipt.parent.mkdir(parents=True, exist_ok=True)
        record = {
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "reason": f"import-{args.kind}",
            "source": args.source.name,
            "packages": versions,
        }
        with args.receipt.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(record, ensure_ascii=False) + "\n")

    packages = ", ".join(f"{name}/{value}" for name, value in versions.items())
    print(f"{packages} ({args.kind})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
